"""Generate document fixtures in code (no binary files committed)."""

from __future__ import annotations

import io
import zipfile

from docx import Document
from fpdf import FPDF


def make_pdf(pages: list[str], *, font_size: int = 12) -> bytes:
    pdf = FPDF()
    pdf.set_font("helvetica", size=font_size)
    for text in pages:
        pdf.add_page()
        if text:
            pdf.multi_cell(0, 6, text)
    return bytes(pdf.output())


def make_blank_pdf(n_pages: int = 3) -> bytes:
    """Pages with no text layer (like a scan whose images are not OCR'd)."""
    pdf = FPDF()
    for _ in range(n_pages):
        pdf.add_page()
        pdf.set_fill_color(200, 200, 200)
        pdf.rect(10, 10, 100, 100, style="F")  # a drawing, no text
    return bytes(pdf.output())


def make_encrypted_pdf() -> bytes:
    from pypdf import PdfReader, PdfWriter

    base = make_pdf(["secret text " * 20])
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(base)))
    w.encrypt("pw", algorithm="RC4-128")
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def make_docx(paragraphs: list[tuple[str, str | None]] | None = None, table: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for text, style in paragraphs or []:
        if style and style.startswith("Heading"):
            doc.add_heading(text, level=int(style.split()[-1]))
        else:
            doc.add_paragraph(text)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_zip(entries: dict[str, bytes], *, compress: bool = True) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED if compress else zipfile.ZIP_STORED) as z:
        for name, data in entries.items():
            z.writestr(name, data)
    return buf.getvalue()


def docx_bomb() -> bytes:
    """A syntactically valid-looking DOCX whose document.xml inflates ~300 MB from a few hundred KB."""
    big = b"<w:t>" + b"A" * (300 * 1024 * 1024)
    return make_zip({"[Content_Types].xml": b"<Types/>", "word/document.xml": big})


PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\rIHDR" + bytes(range(256)) * 4
EXE_BYTES = b"MZ\x90\x00\x03\x00\x00\x00" + b"\x00" * 200 + bytes(range(256))
